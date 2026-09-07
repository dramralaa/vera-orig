"""
# ==============================================================================
# 📤 vera_hungerstation_center/app.py — مركز أتمتة رفع الأصناف الجديدة لفيرا
# ==============================================================================
# تطبيق Streamlit مستقل (نفس فكرة doom_trial_report في ريبو juleb-daily-report):
# مجلد خاص بيه من غير أي pages/ جنبه، عشان ينتشر لوحده على Streamlit Cloud من
# غير ما يجيب أي تاب تاني معاه.
#
# الوظيفة: يجهز ملف/ملفات الرفع اليدوي لأصناف فيرا الجديدة على هنجرستيشن -
# بنفس فورم هنجرستيشن الرسمي بالحرف - لكل فرع على حدة (الروابي / النرجس)،
# مع إمكانية تعديل السعر لكل صنف (أو زيادة نسبة مئوية عامة) قبل التصدير.
#
# ⚠️ ما زال رفع الملف الناتج على Partner Portal خطوة يدوية بمعرفة صاحب
# المشروع - أُثبت هذا الموسم (7 سبتمبر 2026) أن POST /catalog الآلي مرفوض
# للسوق السعودي (501 Not Implemented)، فمفيش بديل آلي للرفع نفسه حاليًا.
#
# مصادر البيانات:
#   - data/reference/vera_master.xlsx: نسخة من ماستر فيرا (Internal Reference,
#     Barcode, Name, Sales Price, Cost, Quantity On Hand, ...) - **سناب شوت
#     ثابت**، مش متصل لحظيًا بجوليب فيرا (مفيش اتصال Odoo مُعد في التطبيق ده
#     لسه) - حدّث الملف يدويًا (استبدل الـ xlsx) لو الأرصدة اتغيّرت (زي بعد
#     الجرد الجاري حاليًا).
#   - data/vera_catalog_al_rawabi.json / vera_catalog_narjis.json: آخر
#     كتالوج حي من هنجرستيشن لكل فرع (بيتحدث تلقائيًا كل 6 ساعات عبر
#     vera_sync_catalog.yml) - يُستخدم لاستبعاد الأصناف الموجودة بالفعل.
#
# النشر على share.streamlit.io: Main file path = vera_hungerstation_center/app.py
# ==============================================================================
"""
import io
import json
import re
from pathlib import Path

import openpyxl
import pandas as pd
import streamlit as st

ACCESS_PASSWORD = "VERA-ADMIN-2026"  # غيّرها هنا لو حبيت

REPO_ROOT = Path(__file__).resolve().parent.parent
MASTER_PATH = REPO_ROOT / "data" / "reference" / "vera_master.xlsx"

BRANCHES = {
    "الروابي (149439)": {"key": "al_rawabi", "catalog": REPO_ROOT / "data" / "vera_catalog_al_rawabi.json"},
    "النرجس (174067)": {"key": "narjis", "catalog": REPO_ROOT / "data" / "vera_catalog_narjis.json"},
}

TEMPLATE_HEADERS = [
    "عنوان المنتج", "السعر (بدون العملة)", "الباركودات(إلزامي)", "SKU",
    "الحد الأقصى لكمية المبيعات", "نشط (1/0)", "الفئة", "صورة (رابط)",
]
MAX_ROWS_PER_FILE = 447

st.set_page_config(page_title="فيرا - مركز رفع الأصناف الجديدة", page_icon="📤", layout="wide")

st.markdown("""
    <style>
    @import url('https://fonts.googleapis.com/css2?family=Cairo:wght@400;700;900&display=swap');
    * { font-family: 'Cairo', sans-serif; }
    .stApp { background-color: #05070a; color: #ffffff; }
    label, div[data-testid="stWidgetLabel"] p { color: #ffffff !important; font-weight: 700 !important; }
    div[data-baseweb="input"] input, div[data-baseweb="select"] > div {
        background-color: #111827 !important; color: #ffffff !important;
    }
    </style>
""", unsafe_allow_html=True)

if "vera_hs_logged_in" not in st.session_state:
    st.session_state.vera_hs_logged_in = False

if not st.session_state.vera_hs_logged_in:
    st.markdown("<h2 style='text-align:center;'>🔐 مركز رفع أصناف فيرا الجديدة - هنجرستيشن</h2>", unsafe_allow_html=True)
    pwd = st.text_input("أدخل رمز الدخول:", type="password")
    if st.button("دخول 🚀"):
        if pwd == ACCESS_PASSWORD:
            st.session_state.vera_hs_logged_in = True
            st.rerun()
        else:
            st.error("رمز الدخول غير صحيح.")
    st.stop()

st.markdown("<h2 style='color:#00d4ff;'>📤 مركز رفع أصناف فيرا الجديدة - هنجرستيشن</h2>", unsafe_allow_html=True)
st.caption("ملف رفع يدوي بنفس فورم هنجرستيشن الرسمي بالحرف - جهّز الأسعار هنا ثم حمّل الملف وارفعه بنفسك على Partner Portal.")
st.markdown("---")


def norm_bc(b):
    return str(b).strip().lstrip("0")


def is_valid_barcode(b):
    return bool(b) and bool(re.fullmatch(r"\d{8,}", str(b).strip()))


def is_valid_name(n):
    s = str(n or "").strip()
    return len(s) >= 4 and not s.isdigit()


@st.cache_data(ttl=600)
def load_master():
    if not MASTER_PATH.exists():
        return pd.DataFrame(), None
    wb = openpyxl.load_workbook(MASTER_PATH, read_only=True, data_only=True)
    ws = wb["Sheet1"]
    rows = []
    seen_refs = set()
    for r in ws.iter_rows(min_row=2, values_only=True):
        ref, barcode, name, price, cost, qty = r[0], r[1], r[2], r[3], r[4], r[5]
        if not ref or str(ref).strip() in seen_refs:
            continue
        seen_refs.add(str(ref).strip())
        rows.append({"ref": str(ref).strip(), "barcode": str(barcode).strip() if barcode else "",
                      "name": name, "price": float(price) if price else 0.0, "qty": int(qty) if qty else 0})
    mtime = MASTER_PATH.stat().st_mtime
    return pd.DataFrame(rows), mtime


@st.cache_data(ttl=600)
def load_existing_barcodes(catalog_path):
    p = Path(catalog_path)
    if not p.exists():
        return set(), set()
    with open(p, encoding="utf-8") as f:
        products = json.load(f)
    exact, normalized = set(), set()
    for prod in products:
        for b in (prod.get("barcodes") or []):
            b = str(b).strip()
            if b:
                exact.add(b)
                normalized.add(norm_bc(b))
    return exact, normalized


def to_excel_bytes(rows):
    buf = io.BytesIO()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(TEMPLATE_HEADERS)
    for r in rows:
        ws.append(r)
    wb.save(buf)
    return buf.getvalue()


master_df, master_mtime = load_master()
if master_df.empty:
    st.error(f"⚠️ ملف الماستر مش موجود أو فاضي: {MASTER_PATH}")
    st.stop()

branch_label = st.selectbox("🏪 اختر الفرع:", list(BRANCHES.keys()))
branch_info = BRANCHES[branch_label]
exact_bc, norm_bc_set = load_existing_barcodes(branch_info["catalog"])

candidates = []
for _, row in master_df.iterrows():
    if row["qty"] <= 0:
        continue
    if not is_valid_barcode(row["barcode"]):
        continue
    if not is_valid_name(row["name"]):
        continue
    if row["barcode"] in exact_bc or norm_bc(row["barcode"]) in norm_bc_set:
        continue
    candidates.append(row)

cand_df = pd.DataFrame(candidates)
st.info(f"📦 عدد الأصناف المرشحة للإضافة في فرع **{branch_label}**: **{len(cand_df)}** "
        f"(رصيد > 0، باركود واسم صالح، ومش موجودة بالفعل في كتالوج الفرع الحالي)")

if cand_df.empty:
    st.stop()

st.markdown("### 💰 التسعير قبل التصدير")
c1, c2 = st.columns(2)
with c1:
    markup_pct = st.number_input("نسبة زيادة عامة على سعر الماستر (%) - تُطبّق على الكل دفعة واحدة:",
                                  min_value=-50.0, max_value=200.0, value=0.0, step=1.0)
with c2:
    default_max_qty = st.number_input("الحد الأقصى الافتراضي لكمية المبيعات لكل صنف (لو الرصيد أكبر منه):",
                                       min_value=1, value=10, step=1)

cand_df = cand_df.copy()
cand_df["price"] = (cand_df["price"] * (1 + markup_pct / 100)).round(2)
cand_df["max_qty"] = cand_df["qty"].clip(upper=default_max_qty)
cand_df["active"] = 1
cand_df["category"] = ""

st.markdown("### ✏️ مراجعة وتعديل الأسعار يدويًا (اختياري) قبل التصدير النهائي")
st.caption("تقدر تعدّل أي سعر أو حد أقصى لصنف بعينه هنا مباشرة قبل ما تصدّر الملف.")
edited = st.data_editor(
    cand_df[["ref", "name", "barcode", "price", "max_qty", "active", "category"]].rename(columns={
        "ref": "SKU", "name": "الاسم", "barcode": "الباركود", "price": "السعر",
        "max_qty": "الحد الأقصى للكمية", "active": "نشط", "category": "الفئة (اختياري)",
    }),
    disabled=["SKU", "الاسم", "الباركود"],
    num_rows="fixed",
    use_container_width=True,
    height=420,
)

st.markdown("### 📥 تصدير ملف الرفع")
total = len(edited)
num_parts = max(1, -(-total // MAX_ROWS_PER_FILE))
st.write(f"إجمالي الأصناف: **{total}** — سيُقسّم إلى **{num_parts}** ملف (حد أقصى {MAX_ROWS_PER_FILE} صف/ملف حسب حد هنجرستيشن).")

for part in range(num_parts):
    chunk = edited.iloc[part * MAX_ROWS_PER_FILE:(part + 1) * MAX_ROWS_PER_FILE]
    rows = [
        [r["الاسم"], r["السعر"], r["الباركود"], r["SKU"], r["الحد الأقصى للكمية"], r["نشط"], r["الفئة (اختياري)"], ""]
        for _, r in chunk.iterrows()
    ]
    xlsx_bytes = to_excel_bytes(rows)
    st.download_button(
        label=f"⬇️ تحميل جزء {part + 1} من {num_parts} ({len(chunk)} صنف)",
        data=xlsx_bytes,
        file_name=f"vera_new_products_{branch_info['key']}_part{part + 1}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        key=f"dl_{branch_info['key']}_{part}",
    )

st.markdown("---")
st.caption("⚠️ الفئة (Category) وتُختار من قائمة هنجرستيشن المحددة مسبقًا - اتركها فاضية هنا وحددها وقت الرفع على Partner Portal، "
           "أو املأها يدويًا في الجدول أعلاه لو عارف القيمة الصحيحة مسبقًا. عمود الصورة يُترك فاضي دايمًا - هنجرستيشن بيربطها تلقائيًا بالباركود.")
